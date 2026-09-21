from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from inscriptions.models import Etudiant
from paiements.models import Paiement
from academique.models import Note, Deliberation, Cours
from documents.models import Document
from django.db.models import Sum, Count, Avg

@staff_member_required
def dashboard_view(request):
    total_etudiants = Etudiant.objects.count()
    etudiants_actifs = Etudiant.objects.filter(statut='actif').count()
    etudiants_preinscrits = Etudiant.objects.filter(statut='preinscrit').count()
    etudiants_suspendus = Etudiant.objects.filter(statut='suspendu').count()
    etudiants_par_filiere = Etudiant.objects.values('filiere').annotate(total=Count('id')).order_by('-total')
    total_encaisse = Paiement.objects.aggregate(total=Sum('montant_paye'))['total'] or 0
    paiements_retard = Paiement.objects.filter(statut='retard').count()
    paiements_soldes = Paiement.objects.filter(statut='solde').count()
    etudiants_dettes = Note.objects.filter(note_finale__lt=10).values('etudiant').distinct().count()
    admis = Deliberation.objects.filter(decision='admis').count()
    ajournes = Deliberation.objects.filter(decision='ajourne').count()
    redoublants = Deliberation.objects.filter(decision='redoublement').count()

    context = {
        'total_etudiants': total_etudiants,
        'etudiants_actifs': etudiants_actifs,
        'etudiants_preinscrits': etudiants_preinscrits,
        'etudiants_suspendus': etudiants_suspendus,
        'etudiants_par_filiere': etudiants_par_filiere,
        'total_encaisse': total_encaisse,
        'paiements_retard': paiements_retard,
        'paiements_soldes': paiements_soldes,
        'etudiants_dettes': etudiants_dettes,
        'admis': admis,
        'ajournes': ajournes,
        'redoublants': redoublants,
    }
    return render(request, 'dashboard/dashboard.html', context)

@staff_member_required
def liste_etudiants(request):
    search = request.GET.get('search', '')
    filiere = request.GET.get('filiere', '')
    statut = request.GET.get('statut', '')

    etudiants = Etudiant.objects.all().order_by('-date_inscription')

    if search:
        etudiants = etudiants.filter(
            nom__icontains=search
        ) | etudiants.filter(
            prenom__icontains=search
        ) | etudiants.filter(
            matricule__icontains=search
        )
    if filiere:
        etudiants = etudiants.filter(filiere=filiere)
    if statut:
        etudiants = etudiants.filter(statut=statut)

    filieres = Etudiant.objects.values_list('filiere', flat=True).distinct()

    context = {
        'etudiants': etudiants,
        'filieres': filieres,
        'search': search,
        'filiere_selected': filiere,
        'statut_selected': statut,
        'total': etudiants.count(),
    }
    return render(request, 'dashboard/liste_etudiants.html', context)

@staff_member_required
def detail_etudiant(request, etudiant_id):
    etudiant = get_object_or_404(Etudiant, id=etudiant_id)
    notes = Note.objects.filter(etudiant=etudiant).select_related('cours')
    paiements = Paiement.objects.filter(etudiant=etudiant)
    documents = Document.objects.filter(etudiant=etudiant)
    deliberation = Deliberation.objects.filter(etudiant=etudiant).last()
    moyenne = notes.aggregate(Avg('note_finale'))['note_finale__avg']

    # Changer statut
    if request.method == 'POST':
        nouveau_statut = request.POST.get('statut')
        if nouveau_statut:
            etudiant.statut = nouveau_statut
            etudiant.save()
            messages.success(request, 'Statut mis à jour !')
            return redirect('detail_etudiant', etudiant_id=etudiant_id)

    context = {
        'etudiant': etudiant,
        'notes': notes,
        'paiements': paiements,
        'documents': documents,
        'deliberation': deliberation,
        'moyenne': moyenne,
    }
    return render(request, 'dashboard/detail_etudiant.html', context)

@staff_member_required
def saisie_notes(request):
    cours_list = Cours.objects.all()
    cours_id = request.GET.get('cours')
    cours_selectionne = None
    etudiants_notes = []

    if cours_id:
        cours_selectionne = get_object_or_404(Cours, id=cours_id)
        etudiants = Etudiant.objects.filter(
            filiere=cours_selectionne.filiere,
            niveau=cours_selectionne.niveau
        )
        for etudiant in etudiants:
            note, created = Note.objects.get_or_create(
                etudiant=etudiant,
                cours=cours_selectionne,
                defaults={'annee_academique': cours_selectionne.annee_academique}
            )
            etudiants_notes.append({'etudiant': etudiant, 'note': note})

    if request.method == 'POST':
        cours_id_post = request.POST.get('cours_id')
        cours_post = get_object_or_404(Cours, id=cours_id_post)
        etudiants = Etudiant.objects.filter(
            filiere=cours_post.filiere,
            niveau=cours_post.niveau
        )
        for etudiant in etudiants:
            note_cc = request.POST.get(f'cc_{etudiant.id}')
            note_exam = request.POST.get(f'exam_{etudiant.id}')
            note, _ = Note.objects.get_or_create(
                etudiant=etudiant,
                cours=cours_post,
                defaults={'annee_academique': cours_post.annee_academique}
            )
            if note_cc:
                note.note_cc = float(note_cc)
            if note_exam:
                note.note_examen = float(note_exam)
            note.save()
        messages.success(request, 'Notes enregistrées avec succès !')
        return redirect(f'/gestion/notes/?cours={cours_id_post}')

    context = {
        'cours_list': cours_list,
        'cours_selectionne': cours_selectionne,
        'etudiants_notes': etudiants_notes,
    }
    return render(request, 'dashboard/saisie_notes.html', context)

@staff_member_required
def liste_paiements(request):
    from paiements.models import Paiement
    from django.db.models import Sum

    paiements = Paiement.objects.all().select_related('etudiant').order_by('-date_paiement')
    
    search = request.GET.get('search', '')
    statut = request.GET.get('statut', '')
    
    if search:
        paiements = paiements.filter(
            etudiant__nom__icontains=search
        ) | paiements.filter(
            etudiant__matricule__icontains=search
        ) | paiements.filter(
            reference__icontains=search
        )
    if statut:
        paiements = paiements.filter(statut=statut)

    total_encaisse = paiements.aggregate(Sum('montant_paye'))['montant_paye__sum'] or 0
    total_reste = sum(p.reste_a_payer for p in paiements)

    context = {
        'paiements': paiements,
        'search': search,
        'statut_selected': statut,
        'total_encaisse': total_encaisse,
        'total_reste': total_reste,
        'total': paiements.count(),
    }
    return render(request, 'dashboard/liste_paiements.html', context)

@staff_member_required
def export_etudiants_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Étudiants IUHEG"

    # Style entête
    header_fill = PatternFill(start_color="1a3a6b", end_color="1a3a6b", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center")

    # Entêtes
    headers = ['Matricule', 'Nom', 'Prénom', 'Email', 'Téléphone',
               'Filière', 'Niveau', 'Statut', 'Année Académique', 'Date Inscription']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align

    # Largeur colonnes
    column_widths = [20, 15, 15, 30, 15, 25, 10, 15, 15, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

    # Données
    etudiants = Etudiant.objects.all().order_by('-date_inscription')
    for row, etudiant in enumerate(etudiants, 2):
        ws.cell(row=row, column=1, value=etudiant.matricule)
        ws.cell(row=row, column=2, value=etudiant.nom)
        ws.cell(row=row, column=3, value=etudiant.prenom)
        ws.cell(row=row, column=4, value=etudiant.email)
        ws.cell(row=row, column=5, value=etudiant.telephone)
        ws.cell(row=row, column=6, value=etudiant.filiere)
        ws.cell(row=row, column=7, value=etudiant.niveau)
        ws.cell(row=row, column=8, value=etudiant.get_statut_display())
        ws.cell(row=row, column=9, value=etudiant.annee_academique)
        ws.cell(row=row, column=10, value=etudiant.date_inscription.strftime('%d/%m/%Y'))

        # Alternance couleurs lignes
        if row % 2 == 0:
            for col in range(1, 11):
                ws.cell(row=row, column=col).fill = PatternFill(
                    start_color="EEF2FF", end_color="EEF2FF", fill_type="solid"
                )

    # Réponse HTTP
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="etudiants_iuheg.xlsx"'
    wb.save(response)
    return response

@staff_member_required
def gestion_deliberations(request):
    from academique.models import Deliberation, Note
    from django.db.models import Avg

    deliberations = Deliberation.objects.all().select_related('etudiant').order_by('-date_deliberation')

    search = request.GET.get('search', '')
    annee = request.GET.get('annee', '')
    decision = request.GET.get('decision', '')

    if search:
        deliberations = deliberations.filter(
            etudiant__nom__icontains=search
        ) | deliberations.filter(
            etudiant__matricule__icontains=search
        )
    if annee:
        deliberations = deliberations.filter(annee_academique=annee)
    if decision:
        deliberations = deliberations.filter(decision=decision)

    annees = Deliberation.objects.values_list('annee_academique', flat=True).distinct()

    context = {
        'deliberations': deliberations,
        'search': search,
        'annee_selected': annee,
        'decision_selected': decision,
        'annees': annees,
        'total': deliberations.count(),
        'admis': deliberations.filter(decision='admis').count(),
        'ajournes': deliberations.filter(decision='ajourne').count(),
        'redoublants': deliberations.filter(decision='redoublement').count(),
    }
    return render(request, 'dashboard/deliberations.html', context)

@staff_member_required
def ajouter_deliberation(request):
    from academique.models import Deliberation
    from django.db.models import Avg
    from academique.models import Note

    if request.method == 'POST':
        etudiant_id = request.POST.get('etudiant_id')
        annee = request.POST.get('annee_academique')
        decision = request.POST.get('decision')
        observations = request.POST.get('observations', '')

        etudiant = get_object_or_404(Etudiant, id=etudiant_id)

        # Calcul moyenne automatique
        notes = Note.objects.filter(etudiant=etudiant, annee_academique=annee)
        moyenne = notes.aggregate(Avg('note_finale'))['note_finale__avg']

        # Créer ou mettre à jour
        delib, created = Deliberation.objects.update_or_create(
            etudiant=etudiant,
            annee_academique=annee,
            defaults={
                'decision': decision,
                'observations': observations,
                'moyenne_generale': moyenne,
            }
        )
        messages.success(request, f'Délibération {"créée" if created else "mise à jour"} pour {etudiant} !')
        return redirect('gestion_deliberations')

    etudiants = Etudiant.objects.filter(statut='actif').order_by('nom')
    context = {'etudiants': etudiants}
    return render(request, 'dashboard/ajouter_deliberation.html', context)

@staff_member_required
def ajouter_paiement(request):
    if request.method == 'POST':
        etudiant_id = request.POST.get('etudiant_id')
        tranche = request.POST.get('tranche')
        montant_total = request.POST.get('montant_total')
        montant_paye = request.POST.get('montant_paye')
        moyen_paiement = request.POST.get('moyen_paiement')
        statut = request.POST.get('statut')
        date_echeance = request.POST.get('date_echeance') or None

        etudiant = get_object_or_404(Etudiant, id=etudiant_id)

        from paiements.models import Paiement
        paiement = Paiement.objects.create(
            etudiant=etudiant,
            tranche=tranche,
            montant_total=montant_total,
            montant_paye=montant_paye,
            moyen_paiement=moyen_paiement,
            statut=statut,
            date_echeance=date_echeance
        )
        messages.success(request, f'Paiement {paiement.reference} enregistré pour {etudiant} !')
        return redirect('liste_paiements')

    etudiants = Etudiant.objects.filter(statut__in=['actif', 'inscrit']).order_by('nom')
    context = {'etudiants': etudiants}
    return render(request, 'dashboard/ajouter_paiement.html', context)

@staff_member_required
def gestion_cours(request):
    from cours.models import Cours
    cours_list = Cours.objects.all().order_by('filiere', 'niveau', 'code')
    filieres = Cours.objects.values_list('filiere', flat=True).distinct()

    filiere = request.GET.get('filiere', '')
    if filiere:
        cours_list = cours_list.filter(filiere=filiere)

    context = {
        'cours_list': cours_list,
        'filieres': filieres,
        'filiere_selected': filiere,
        'total': cours_list.count(),
    }
    return render(request, 'dashboard/cours.html', context)

@staff_member_required
def ajouter_cours(request):
    from cours.models import Cours
    if request.method == 'POST':
        Cours.objects.create(
            nom=request.POST.get('nom'),
            code=request.POST.get('code'),
            credits=request.POST.get('credits', 3),
            coefficient=request.POST.get('coefficient', 1),
            filiere=request.POST.get('filiere'),
            niveau=request.POST.get('niveau'),
            annee_academique=request.POST.get('annee_academique'),
            semestre=request.POST.get('semestre'),
        )
        messages.success(request, 'Cours ajouté avec succès !')
        return redirect('gestion_cours')
    return render(request, 'dashboard/ajouter_cours.html')

@staff_member_required
def reinscription(request):
    if request.method == 'POST':
        etudiant_id = request.POST.get('etudiant_id')
        nouvelle_annee = request.POST.get('nouvelle_annee')
        nouveau_niveau = request.POST.get('nouveau_niveau')

        etudiant_source = get_object_or_404(Etudiant, id=etudiant_id)

        # Vérifier si une réinscription existe déjà pour cette année
        existe = Etudiant.objects.filter(
            email=etudiant_source.email,
            annee_academique=nouvelle_annee
        ).exists()

        if existe:
            messages.error(request, f'{etudiant_source.nom} {etudiant_source.prenom} est déjà inscrit(e) pour {nouvelle_annee}.')
        else:
            # Créer un nouveau dossier étudiant pour la nouvelle année
            # en conservant toutes les informations personnelles
            nouvel_etudiant = Etudiant.objects.create(
                nom=etudiant_source.nom,
                prenom=etudiant_source.prenom,
                date_naissance=etudiant_source.date_naissance,
                lieu_naissance=etudiant_source.lieu_naissance,
                sexe=etudiant_source.sexe,
                nationalite=etudiant_source.nationalite,
                telephone=etudiant_source.telephone,
                email=etudiant_source.email,
                adresse=etudiant_source.adresse,
                filiere=etudiant_source.filiere,
                niveau=nouveau_niveau,
                annee_academique=nouvelle_annee,
                statut='inscrit',
            )
            messages.success(
                request,
                f'{nouvel_etudiant.nom} {nouvel_etudiant.prenom} réinscrit(e) pour {nouvelle_annee} '
                f'avec le nouveau matricule {nouvel_etudiant.matricule} !'
            )
        return redirect('reinscription')

    # Étudiants actifs de l'année précédente, éligibles à la réinscription
    annee_active = request.GET.get('annee_source', '')
    etudiants = Etudiant.objects.filter(statut='actif').order_by('nom')
    if annee_active:
        etudiants = etudiants.filter(annee_academique=annee_active)

    annees_existantes = Etudiant.objects.values_list('annee_academique', flat=True).distinct().order_by('-annee_academique')

    context = {
        'etudiants': etudiants,
        'annees_existantes': annees_existantes,
        'annee_active': annee_active,
    }
    return render(request, 'dashboard/reinscription.html', context)

@staff_member_required
def historique_etudiant(request, email):
    """Affiche tout l'historique d'un étudiant à travers les années académiques"""
    dossiers = Etudiant.objects.filter(email=email).order_by('annee_academique')
    context = {'dossiers': dossiers, 'email': email}
    return render(request, 'dashboard/historique_etudiant.html', context)

@staff_member_required
def gestion_utilisateurs(request):
    from django.contrib.auth.models import User, Group

    users = User.objects.all().order_by('username').prefetch_related('groups')
    groupes = Group.objects.all()

    context = {
        'users': users,
        'groupes': groupes,
        'total': users.count(),
    }
    return render(request, 'dashboard/utilisateurs.html', context)

@staff_member_required
def ajouter_utilisateur(request):
    from django.contrib.auth.models import User, Group

    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        groupe_id = request.POST.get('groupe')
        prenom = request.POST.get('prenom')
        nom = request.POST.get('nom')

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Ce nom d\'utilisateur existe déjà.')
        else:
            user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
                first_name=prenom,
                last_name=nom,
                is_staff=True
            )
            if groupe_id:
                groupe = Group.objects.get(id=groupe_id)
                user.groups.add(groupe)
            messages.success(request, f'Utilisateur {username} créé avec succès !')
            return redirect('gestion_utilisateurs')

    groupes = Group.objects.all()
    return render(request, 'dashboard/ajouter_utilisateur.html', {'groupes': groupes})

@staff_member_required
def historique_paiements(request):
    from paiements.models import Paiement

    paiements = Paiement.objects.all().select_related('etudiant').order_by('-date_paiement')

    search = request.GET.get('search', '')
    statut = request.GET.get('statut', '')
    moyen = request.GET.get('moyen', '')

    if search:
        paiements = paiements.filter(
            etudiant__nom__icontains=search
        ) | paiements.filter(
            etudiant__matricule__icontains=search
        ) | paiements.filter(
            reference__icontains=search
        )
    if statut:
        paiements = paiements.filter(statut=statut)
    if moyen:
        paiements = paiements.filter(moyen_paiement=moyen)

    from django.db.models import Sum
    total_encaisse = paiements.aggregate(Sum('montant_paye'))['montant_paye__sum'] or 0

    context = {
        'paiements': paiements,
        'search': search,
        'statut_selected': statut,
        'moyen_selected': moyen,
        'total': paiements.count(),
        'total_encaisse': total_encaisse,
    }
    return render(request, 'dashboard/historique_paiements.html', context)

def login_admin(request):
    from django.contrib.auth import authenticate, login as auth_login

    if request.user.is_authenticated and request.user.is_staff:
        return redirect('dashboard')

    error = None
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None and user.is_staff:
            auth_login(request, user)
            return redirect('dashboard')
        else:
            error = "Identifiants incorrects ou accès non autorisé."

    return render(request, 'dashboard/login_admin.html', {'error': error})

@staff_member_required
def gestion_documents(request):
    from documents.models import Document
    documents = Document.objects.all().select_related('etudiant').order_by('-date_upload')
    search = request.GET.get('search', '')
    statut = request.GET.get('statut', '')
    if search:
        documents = documents.filter(etudiant__nom__icontains=search) | documents.filter(etudiant__matricule__icontains=search)
    if statut:
        documents = documents.filter(statut=statut)
    context = {
        'documents': documents,
        'search': search,
        'statut_selected': statut,
        'total': documents.count(),
        'en_attente': documents.filter(statut='en_attente').count(),
        'valides': documents.filter(statut='valide').count(),
        'rejetes': documents.filter(statut='rejete').count(),
    }
    return render(request, 'dashboard/documents.html', context)

@staff_member_required
def valider_document(request, doc_id):
    from documents.models import Document
    doc = get_object_or_404(Document, id=doc_id)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'valider':
            doc.statut = 'valide'
            messages.success(request, f'Document validé !')
        elif action == 'rejeter':
            doc.statut = 'rejete'
            messages.warning(request, f'Document rejeté.')
        doc.save()
    return redirect('gestion_documents')
